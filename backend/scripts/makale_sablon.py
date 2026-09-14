"""Makaleyi CDEJ'in KENDİ .docx şablonunun içine yerleştirir.

⚠ NEDEN ŞABLONUN İÇİNE
İlk denemede şablonun biçim kurallarına uyan YENİ bir belge üretilmişti.
Kullanıcı haklı olarak itiraz etti: derginin editörü "şablonu da
hazırla" demişti ve yayımlanmış örnek makale, ilk sayfanın dergi
başlığı, öz kutusu ve künye sütunu ile birlikte geldiğini gösteriyor.
Bu betik şablon dosyasını açıp İÇERİĞİ DEĞİŞTİRİR; başlık görselleri,
öz kutusu ve sayfa düzeni şablondan gelir.

⚠ ŞABLONUN YAPISI (incelenerek çıkarıldı)
Şablonda 8 metin kutusu var, her biri iki kez geçiyor (Word'ün eski
sürümleri için `mc:AlternateContent` yedeği). Bu yüzden her kutu
İKİ KEZ doldurulmalı; yalnızca birini doldurmak, belgeyi açan Word
sürümüne göre farklı içerik göstermesine yol açar.

  kutu 0/1   : dergi adı (değişmez)
  kutu 2/3   : cilt, sayı, DOI  -> DERGİ DOLDURUR, boş bırakılıyor
  kutu 4/5   : e-ISSN ve adres (değişmez)
  kutu 6/7   : "Cite as" künyesi
  kutu 8/9   : Received / Revision / Accepted / Online -> DERGİ DOLDURUR
  kutu 10/11 : Öz ve anahtar kelimeler
  kutu 12/13 : kurum bilgileri ve yazışma yazarı
  kutu 14/15 : "Research Article"

Gövde ise normal paragraflarda; şablonun örnek metni silinip yerine
Markdown'dan üretilen içerik konuyor.

⚠ GİRDİLERİ BU DEPODA DEĞİL (12.09.2026)
Bu betik depoda duruyor ama beslendiği dosyalar durmuyor:
`docs/report/makale/` `.gitignore` içinde. Sebep, dergiye sunulacak
metnin ve kurum şablonunun genel bir depoya girmemesi. Depoyu
klonlayan biri bu betiği olduğu gibi çalıştıramaz; betik burada
yöntemi belgelemek için duruyor.

Kullanım:
    uv run --with python-docx python scripts/makale_sablon.py --dil tr
    uv run --with python-docx python scripts/makale_sablon.py --dil en
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]

KOK = Path(__file__).resolve().parents[2]
SABLON = Path.home() / "Downloads" / "templateCDEJ_ENG.docx"
SEKIL = KOK / "docs" / "report" / "diagrams"
CIKTI_DIZIN = KOK / "docs" / "report" / "makale"

# ⚠ SÜRÜMLER (13.09.2026) — eski sürümün dosyalarının ÜSTÜNE YAZILMAZ.
# v1: danışmana 12.09'da gönderilen hâl. v2: danışman geri bildirimi —
# başlık "Yirmi Kameralı" → "Çoklu Kamera ile" ve şekiller Elsevier
# tarzında yeniden çizildi (scripts/sekiller_elsevier.py). v2'de şekiller
# DİLE GÖRE ayrı: İngilizce makale İngilizce şekil kullanır; v1'de iki
# dil de aynı Türkçe şekilleri kullanıyordu.
SURUMLER = {
    "v1": {
        "kaynak": {"tr": CIKTI_DIZIN / "05-makale-tr-final.md",
                   "en": CIKTI_DIZIN / "06-makale-en-final.md"},
        "cikti": {"tr": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-TR.docx",
                  "en": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-ENG.docx"},
        "sekil": lambda dil, no: SEKIL / {"1": "01-veri-akisi.png",
                                          "2": "02-kademeli-isleme.png",
                                          "3": "03-kimlik-dogrulama.png"}.get(no, ""),
        "sekil_cm": 15.5,
    },
    "v2": {
        "kaynak": {"tr": CIKTI_DIZIN / "07-makale-tr-v2.md",
                   "en": CIKTI_DIZIN / "08-makale-en-v2.md"},
        "cikti": {"tr": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-TR-v2.docx",
                  "en": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-ENG-v2.docx"},
        "sekil": lambda dil, no: SEKIL / "elsevier" / dil / {
            "1": "sekil-1-veri-akisi.png", "2": "sekil-2-kademeli-isleme.png",
            "3": "sekil-3-kimlik-dogrulama.png"}.get(no, ""),
        # Şablonun metin genişliği 17,99 cm; şekil tam metin genişliğinde.
        "sekil_cm": 17.8,
    },
}
# v3 (14.09): beyanlardaki intihal cümlesi çıkarıldı — tarama yapılmadan
# şablonun koşullu örnek cümlesi yazılmıştı; taramayı dergi yapıyor.
# v3-kor: çift kör hakemlik için yazar kimliği çıkarılmış kopya (yalnız EN).
SURUMLER["v3"] = {**SURUMLER["v2"],
                  "kaynak": {"tr": CIKTI_DIZIN / "09-makale-tr-v3.md",
                             "en": CIKTI_DIZIN / "10-makale-en-v3.md"},
                  "cikti": {"tr": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-TR-v3.docx",
                            "en": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-ENG-v3.docx"}}
SURUMLER["v3-kor"] = {**SURUMLER["v2"],
                      "kaynak": {"en": CIKTI_DIZIN / "10-makale-en-v3-kor-hakem.md"},
                      "cikti": {"en": CIKTI_DIZIN
                                / "SENTINEL-CDEJ-sablon-ENG-v3-kor-hakem.docx"},
                      "kor": True}
# v4 (14.09): yapay zekâ beyanına model sürümü eklendi (Claude Opus 5) —
# derginin YZ politikası "tam ad ve sürüm numarası" istiyor. Başka değişiklik yok.
SURUMLER["v4"] = {**SURUMLER["v2"],
                  "kaynak": {"tr": CIKTI_DIZIN / "11-makale-tr-v4.md",
                             "en": CIKTI_DIZIN / "12-makale-en-v4.md"},
                  "cikti": {"tr": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-TR-v4.docx",
                            "en": CIKTI_DIZIN / "SENTINEL-CDEJ-sablon-ENG-v4.docx"}}
SURUMLER["v4-kor"] = {**SURUMLER["v3-kor"],
                      "kaynak": {"en": CIKTI_DIZIN / "12-makale-en-v4-kor-hakem.md"},
                      "cikti": {"en": CIKTI_DIZIN
                                / "SENTINEL-CDEJ-sablon-ENG-v4-kor-hakem.docx"}}
KOR = "[Removed for blind review]"
YAZI = "Cambria"


# ──────────────────────────────────────────────────────────────
#  Metin kutusu doldurma
# ──────────────────────────────────────────────────────────────
def _kutular(belge):
    """Belgedeki tüm `w:txbxContent` ögelerini belge sırasıyla verir."""
    ad = qn("w:txbxContent")
    return belge.element.body.iter(ad)


def _kutu_yaz(kutu, paragraflar: list[list[tuple[str, bool]]]) -> None:
    """Kutunun içeriğini yeniden yazar; biçimi mevcut ilk koşudan alır.

    `paragraflar`: her biri (metin, kalın_mı) ikililerinden oluşan liste.

    ⚠ Biçim, kutudaki İLK koşunun `rPr`'sinden kopyalanıyor. Böylece
    yazı tipi, punto ve renk şablonun kendi tanımından geliyor; burada
    yeniden tanımlanmıyor.
    """
    mevcut = list(kutu.findall(qn("w:p")))
    if not mevcut:
        return
    ornek_p = mevcut[0]
    ppr = ornek_p.find(qn("w:pPr"))
    ornek_r = ornek_p.find(qn("w:r"))
    rpr = ornek_r.find(qn("w:rPr")) if ornek_r is not None else None

    for p in mevcut:
        kutu.remove(p)

    for satirlar in paragraflar:
        yeni_p = copy.deepcopy(ornek_p)
        for c in list(yeni_p):
            if c.tag != qn("w:pPr"):
                yeni_p.remove(c)
        if ppr is None and yeni_p.find(qn("w:pPr")) is not None:
            yeni_p.remove(yeni_p.find(qn("w:pPr")))
        for metin, kalin in satirlar:
            r = yeni_p.makeelement(qn("w:r"), {})
            if rpr is not None:
                yeni_rpr = copy.deepcopy(rpr)
                b = yeni_rpr.find(qn("w:b"))
                if kalin and b is None:
                    yeni_rpr.append(yeni_rpr.makeelement(qn("w:b"), {}))
                elif not kalin and b is not None:
                    yeni_rpr.remove(b)
                r.append(yeni_rpr)
            t = r.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
            t.text = metin
            r.append(t)
            yeni_p.append(r)
        kutu.append(yeni_p)


def _derin_degistir(belge, baslangic: str, yeni: str) -> int:
    """Metni, nerede olursa olsun (tablo/metin kutusu içinde dâhil) değiştirir.

    Paragrafın koşulları birleştirilip ilk koşuya yazılıyor, diğerleri
    boşaltılıyor. Böylece ilk koşunun biçimi korunuyor ve metin
    koşulara bölünmüş olsa da doğru sonuç elde ediliyor.
    """
    # ⚠⚠ `.//w:t` KULLANILMAZ. Sayfa 1'deki kayan metin kutuları (öz,
    # künye, "Cite as") aynı paragrafa bağlı çizimlerin içinde duruyor.
    # Derinlemesine arama, bir kutuyu değiştirirken AYNI paragrafa bağlı
    # diğer kutuların metnini de siliyordu; ilk denemede öz ve künye
    # kutuları bu yüzden boş çıktı. Yalnızca paragrafın DOĞRUDAN
    # koşuları taranıyor.
    sayi = 0
    for p in belge.element.body.iter(qn("w:p")):
        tler = [t for r in p.findall(qn("w:r")) for t in r.findall(qn("w:t"))]
        if not tler:
            continue
        tam = "".join(t.text or "" for t in tler)
        if baslangic not in tam:
            continue
        tler[0].text = yeni
        tler[0].set(qn("xml:space"), "preserve")
        for t in tler[1:]:
            t.text = ""
        sayi += 1
    return sayi


# ──────────────────────────────────────────────────────────────
#  Gövde
# ──────────────────────────────────────────────────────────────
def _tipi(calisma, punto, kalin=False, italik=False):
    calisma.font.name = YAZI
    calisma.font.size = Pt(punto)
    calisma.font.bold = kalin
    calisma.font.italic = italik
    rf = calisma._element.rPr.rFonts
    rf.set(qn("w:eastAsia"), YAZI)
    rf.set(qn("w:cs"), YAZI)


def _zengin(par, metin, punto, kalin=False, italik=False):
    for parca in re.split(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)", metin):
        if not parca:
            continue
        k, i = kalin, italik
        if parca.startswith("**") and parca.endswith("**"):
            parca, k = parca[2:-2], True
        elif parca.startswith("*") and parca.endswith("*"):
            parca, i = parca[1:-1], True
        elif parca.startswith("`") and parca.endswith("`"):
            parca = parca[1:-1]
        _tipi(par.add_run(parca), punto, k, i)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dil", choices=("tr", "en"), default="tr")
    ap.add_argument("--surum", choices=tuple(SURUMLER), default="v1")
    args = ap.parse_args()
    dil = args.dil
    surum = SURUMLER[args.surum]

    if not SABLON.is_file():
        print(f"❌ Şablon bulunamadı: {SABLON}")
        return 1
    if dil not in surum["kaynak"]:
        print(f"❌ {args.surum} sürümünde '{dil}' dili yok")
        return 1
    kaynak = surum["kaynak"][dil]
    if not kaynak.is_file():
        print(f"❌ Kaynak metin yok: {kaynak}")
        return 1

    ham = kaynak.read_text(encoding="utf-8")
    belge = Document(str(SABLON))

    # ── Markdown'dan üstbilgi alanlarını ayıkla ──
    baslik = re.search(r"^# (.+)$", ham, re.M).group(1).strip()
    oz_m = re.search(r"^## (?:ÖZ|ABSTRACT)\s*\n+(.+?)\n+\*\*(?:Anahtar Kelimeler|Keywords):\*\*(.+?)$",
                     ham, re.M | re.S)
    oz_metin = oz_m.group(1).strip()
    anahtar = oz_m.group(2).strip()

    oz_etiket = "Öz: " if dil == "tr" else "Abstract: "
    anahtar_etiket = "Anahtar Kelimeler: " if dil == "tr" else "Keywords: "

    kurum = ("1 Düzce Üniversitesi, Mühendislik Fakültesi, Bilgisayar Mühendisliği "
             "Bölümü, Düzce, Türkiye.")
    ror = "ror.org/04175wc52"
    yazisma = "*Yazışma yazarı:" if dil == "tr" else "*Corresponding author:"
    adsoyad = "Ömer Faruk Kanat"
    eposta = "omerfk0121@gmail.com"
    # ⚠ "ORCID" BÜYÜK HARF yazılır. 12.09 lint düzeltmesinde büyük harfli
    # değişken adları küçültülürken regex string'in İÇİNDEKİ "ORCID" yazısını
    # da küçültmüştü; belgeye "orcid:" olarak girdi. v1/v2 metin karşılaştırması
    # (yalnızca başlık farkı beklenirken) bunu 13.09'da yakaladı.
    orcid = "ORCID: 0009-0006-9229-6217"

    kor = surum.get("kor", False)
    if kor:
        # ⚠ Çift kör hakemlik: kimlik açan HER alan boşaltılıyor. Kutu
        # yapısı korunuyor ki şablon düzeni bozulmasın.
        kurum, ror, yazisma, eposta, orcid = KOR, "", "", "", ""
        adsoyad = "[Author name removed for blind review]"

    cite = (f"Cite as: {'[Removed for blind review]' if kor else 'Kanat, Ö. F.'} (202x). {baslik} "
            f"Cyber Security and Digital Economy Journal, vol(issue), xx-xx.")
    lisans = ("This is an open access paper distributed under the terms and conditions "
              "of the Creative Commons Attribution-NonCommercial 4.0 International License.")

    # ── Metin kutularını doldur (her biri İKİ KEZ geçiyor) ──
    kutular = list(_kutular(belge))
    print(f"metin kutusu: {len(kutular)}")
    icerik = {
        6: [[(cite, False)], [(lisans, False)]],
        7: [[(cite, False)], [(lisans, False)]],
        10: [[(oz_etiket, True), (oz_metin, False)], [("", False)],
             [(anahtar_etiket, True), (anahtar, False)]],
        11: [[(oz_etiket, True), (oz_metin, False)], [("", False)],
             [(anahtar_etiket, True), (anahtar, False)]],
        12: [[(kurum, False)], [(ror, False)], [("", False)],
             [(yazisma, False)], [(adsoyad, False)], [(eposta, False)], [(orcid, False)]],
        13: [[(kurum, False)], [(ror, False)], [("", False)],
             [(yazisma, False)], [(adsoyad, False)], [(eposta, False)], [(orcid, False)]],
    }
    for indeks, paragraflar in icerik.items():
        if indeks < len(kutular):
            _kutu_yaz(kutular[indeks], paragraflar)

    # ⚠ "Cite as" kutusunun içinde bir TABLO var ve metin onun içinde.
    # Doğrudan çocuklara bakan `_kutu_yaz` bu kutuyu boş geçiyordu;
    # metin derinlemesine aranıp değiştiriliyor.
    _derin_degistir(belge, "Cite as: Surname", cite)
    _derin_degistir(belge, "This is an open access paper", lisans)

    # ── Başlık ve yazar ──
    for p in belge.paragraphs:
        if p.style.name == "Heading 1" and "Article Title" in p.text:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            _tipi(p.add_run(baslik), 12, kalin=True)
        if p.text.strip().startswith("Beytullah Çıtır"):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            _tipi(p.add_run(adsoyad if kor else f"{adsoyad}¹*"), 11, kalin=False)

    # ── Şablonun örnek gövdesini sil ──
    govde = belge.element.body
    ogeler = list(govde.iterchildren())
    baslangic = None
    for i, c in enumerate(ogeler):
        if c.tag == qn("w:p") and Paragraph(c, belge).text.strip().startswith("I. INTRODUCTION"):
            baslangic = i
            break
    if baslangic is None:
        print("❌ Şablonda 'I. INTRODUCTION' bulunamadı")
        return 1
    son = len(ogeler)
    for c in ogeler[baslangic:son]:
        if c.tag == qn("w:sectPr"):
            continue
        govde.remove(c)
    print(f"şablon gövdesinden silinen öge: {son - baslangic}")

    # ── Kendi gövdemizi ekle ──
    # ⚠ Türkçe makalede gövde "ABSTRACT" ile başlıyor. Şablon kuralı:
    # "Türkçe yazılan makaleler Türkçe öz (en fazla 250 kelime) VE
    # İngilizce abstract içermelidir." Türkçe öz sayfa 1'deki kutuda,
    # İngilizce abstract ise gövdenin başında yer alıyor.
    # İngilizce makalede abstract zaten kutuda; gövde I. bölümle başlar.
    _govde_yaz(belge, ham, "## ABSTRACT" if dil == "tr" else None,
               sekil_yolu=lambda no: surum["sekil"](dil, no),
               sekil_cm=surum["sekil_cm"])

    if kor:
        # Word dosya özellikleri de yazarı ele verebilir (Dosya > Bilgi)
        cp = belge.core_properties
        cp.author = cp.last_modified_by = ""
        cp.title = cp.subject = cp.keywords = cp.comments = ""

    CIKTI_DIZIN.mkdir(parents=True, exist_ok=True)
    cikti = surum["cikti"][dil]
    belge.save(str(cikti))
    print(f"yazıldı: {cikti.relative_to(KOK)}")
    print(f"paragraf: {len(belge.paragraphs)} · tablo: {len(belge.tables)}")
    return 0


def _govde_yaz(belge, ham: str, ilk_baslik: str | None = None,
               sekil_yolu=None, sekil_cm: float = 15.5) -> None:
    """Markdown gövdesini belgeye ekler.

    `ilk_baslik` verilirse gövde o başlıktan, verilmezse "## I." ile
    başlar. Türkçe makalede İngilizce abstract gövdenin ilk ögesidir.
    """
    satirlar = ham.splitlines()
    if ilk_baslik:
        bas = next(i for i, s in enumerate(satirlar) if s.strip() == ilk_baslik)
    else:
        bas = next(i for i, s in enumerate(satirlar)
                   if re.match(r"^## (I\.|I\s)", s.strip()))
    i, n = bas, len(satirlar)

    def par(metin="", punto=10.0, kalin=False, italik=False,
            hiza=WD_ALIGN_PARAGRAPH.JUSTIFY, sonra=6, once=0):
        p = belge.add_paragraph()
        p.alignment = hiza
        p.paragraph_format.space_after = Pt(sonra)
        p.paragraph_format.space_before = Pt(once)
        p.paragraph_format.line_spacing = 1.0
        if metin:
            _zengin(p, metin, punto, kalin, italik)
        return p

    def md_tablo(s):
        return s.strip().startswith("|") and s.strip().endswith("|")

    while i < n:
        s = satirlar[i]
        k = s.strip()
        if not k or k == "---" or k.startswith(">"):
            i += 1
            continue

        if md_tablo(k):
            blok = []
            while i < n and md_tablo(satirlar[i].strip()):
                blok.append([h.strip() for h in satirlar[i].strip().strip("|").split("|")])
                i += 1
            t = belge.add_table(rows=1, cols=len(blok[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for h, b in zip(t.rows[0].cells, blok[0], strict=False):
                h.text = ""
                hp = h.paragraphs[0]
                hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                hp.paragraph_format.space_after = Pt(2)
                _zengin(hp, b, 8.5, kalin=True)
            for satir in blok[2:]:
                hucre = t.add_row().cells
                for j, (h, d) in enumerate(zip(hucre, satir, strict=False)):
                    h.text = ""
                    hp = h.paragraphs[0]
                    hp.alignment = (WD_ALIGN_PARAGRAPH.LEFT if j == 0
                                    else WD_ALIGN_PARAGRAPH.CENTER)
                    hp.paragraph_format.space_after = Pt(2)
                    _zengin(hp, d, 8.5)
            par("", punto=6, sonra=6)
            continue

        if k.startswith("## "):
            par(k[3:], punto=10.5, kalin=True, hiza=WD_ALIGN_PARAGRAPH.CENTER,
                once=12, sonra=8)
        elif k.startswith("### "):
            par(k[4:], punto=10, kalin=True, italik=True,
                hiza=WD_ALIGN_PARAGRAPH.LEFT, once=8, sonra=4)
        elif k.startswith("#### "):
            par(k[5:], punto=10, italik=True, hiza=WD_ALIGN_PARAGRAPH.LEFT,
                once=6, sonra=4)
        elif k.startswith("**[ŞEKİL") or k.startswith("**[FIGURE"):
            no = re.search(r"(?:ŞEKİL|FIGURE) (\d+)", k)
            yol = sekil_yolu(no.group(1)) if (no and sekil_yolu) else None
            p = belge.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if yol and yol.is_file():
                p.add_run().add_picture(str(yol), width=Cm(sekil_cm))
            else:
                print(f"  ⚠ şekil bulunamadı: {k} → {yol}")
                _tipi(p.add_run(k.strip("*[]")), 9, italik=True)
        elif k.startswith("*Şekil ") or k.startswith("*Figure "):
            par(k.strip("*"), punto=9, italik=True,
                hiza=WD_ALIGN_PARAGRAPH.CENTER, sonra=10)
        elif k.startswith("*Tablo ") or k.startswith("*Table "):
            par(k.strip("*"), punto=9, italik=True,
                hiza=WD_ALIGN_PARAGRAPH.CENTER, once=8, sonra=3)
        else:
            par(k, punto=10)
        i += 1


if __name__ == "__main__":
    raise SystemExit(main())
